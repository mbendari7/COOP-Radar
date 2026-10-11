"""
Resume test fixtures — in-memory PDF and DOCX bytes with skill keywords.
"""

import io

import pytest


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """
    Minimal valid PDF containing resume-like text with skill keywords.
    Generated programmatically so tests have no binary file dependency.
    """
    content = (
        "John Doe\n"
        "Software Developer\n"
        "Skills: Python, Django, PostgreSQL, Docker, REST APIs\n"
        "Experience: 2 years backend development with Python and Django.\n"
        "Built REST APIs using FastAPI and PostgreSQL.\n"
        "Deployed containerised services with Docker on AWS.\n"
    )
    return _build_minimal_pdf(content)


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """
    Minimal valid DOCX containing resume-like text with skill keywords.
    """
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_paragraph("Jane Smith")
    doc.add_paragraph("Software Engineer Co-op")
    doc.add_paragraph(
        "Skills: Python, React, TypeScript, PostgreSQL, Docker, Kubernetes, AWS"
    )
    doc.add_paragraph(
        "Developed full-stack web applications using React and FastAPI. "
        "Managed PostgreSQL databases and deployed with Docker."
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_minimal_pdf(text: str) -> bytes:
    """
    Build a minimal valid single-page PDF with embedded text.
    This bypasses pypdf's lack of text-writing API for blank pages.
    """
    lines = text.strip().splitlines()
    stream_lines = []
    y = 750
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_lines.append(f"BT /F1 12 Tf 50 {y} Td ({safe}) Tj ET")
        y -= 20

    stream = "\n".join(stream_lines)
    stream_bytes = stream.encode("latin-1")

    objects: list[bytes] = []

    def obj(n: int, content: str) -> bytes:
        return f"{n} 0 obj\n{content}\nendobj\n".encode()

    objects.append(obj(1, "<< /Type /Catalog /Pages 2 0 R >>"))
    objects.append(obj(2, "<< /Type /Pages /Kids [3 0 R] /Count 1 >>"))
    objects.append(
        obj(3, "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
               "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>")
    )
    objects.append(obj(4, f"<< /Length {len(stream_bytes)} >>\nstream\n" + stream + "\nendstream"))
    objects.append(obj(5, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"))

    header = b"%PDF-1.4\n"
    body = b"".join(objects)
    xref_offset = len(header) + len(body)

    xref = f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    pos = len(header)
    for o in objects:
        xref += f"{pos:010d} 00000 n \n"
        pos += len(o)

    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    )

    return header + body + xref.encode() + trailer.encode()
